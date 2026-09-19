import fitz  # PyMuPDF
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
import io


def extract_pdf(file_bytes: bytes) -> str:
    """Extracts text from a PDF, inserting [PAGE X] markers."""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    extracted_text = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text()
        extracted_text.append(f"[PAGE {page_num + 1}]\n{text}")

    return "\n\n".join(extracted_text)


def _docx_table_to_markdown(table: Table) -> str:
    """Converts a python-docx Table into a Markdown table string."""
    rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""

    header = rows[0]
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join(["---"] * len(header)) + "|"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def _iter_block_items(doc: Document):
    """
    Yields each paragraph and table in a docx in the order they actually
    appear in the document body, so tables aren't skipped or reordered.
    python-docx's `doc.paragraphs` / `doc.tables` expose these as two
    separate flat lists with no positional relationship to each other.
    """
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def extract_docx(file_bytes: bytes) -> str:
    """
    Extracts text from a Word document, preserving paragraph/table order.
    Tables are rendered as Markdown so row/column structure survives instead
    of being silently dropped (doc.paragraphs alone skips tables entirely).
    """
    doc = Document(io.BytesIO(file_bytes))
    parts = []

    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text)
        elif isinstance(block, Table):
            markdown_table = _docx_table_to_markdown(block)
            if markdown_table:
                parts.append(markdown_table)

    return "\n".join(parts)


def concatenate_files(files: list[dict]) -> str:
    """
    Joins multiple files into one large text blob.
    'files' should be a list of dicts: [{"filename": "contract.pdf", "text": "..."}]
    """
    if not files:
        return ""

    if len(files) == 1:
        # If there's only one file, we don't need FILE markers
        return files[0]["text"]

    combined = []
    for f in files:
        combined.append(f"[FILE: {f['filename']}]\n{f['text']}")

    return "\n\n---\n\n".join(combined)
